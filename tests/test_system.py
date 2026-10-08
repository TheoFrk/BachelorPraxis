import csv
import io
import sqlite3
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from securities_eye import store, context
from securities_eye.app import create_app

TOKEN = "test-token-with-at-least-32-characters"
ORIGIN = {"origin": "http://localhost:8000"}


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "app.sqlite"
    store.initialize(path)
    return path


@pytest.fixture
def client(db, monkeypatch):
    monkeypatch.setenv("SE_SECURE_COOKIE", "0")
    with TestClient(create_app(db, TOKEN)) as client:
        yield client


def prices(symbol="SAP.DE", count=25):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["symbol", "date", "close", "name", "currency"])
    for i in range(count):
        writer.writerow(
            [
                symbol,
                (date(2025, 1, 1) + timedelta(days=i)).isoformat(),
                100 + i,
                "SAP",
                "EUR",
            ]
        )
    return out.getvalue()


def login(client):
    assert (
        client.post("/api/session", json={"token": TOKEN}, headers=ORIGIN).status_code
        == 200
    )


def test_auth_origin_and_cookie_security(client):
    assert client.get("/health/ready").status_code == 200
    assert client.get("/api/market").status_code == 401
    assert client.post("/api/session", json={"token": TOKEN}).status_code == 403
    assert (
        client.post(
            "/api/session",
            json={"token": TOKEN},
            headers={"origin": "https://evil.example"},
        ).status_code
        == 403
    )
    response = client.post("/api/session", json={"token": TOKEN}, headers=ORIGIN)
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=strict" in response.headers["set-cookie"]
    assert client.get("/api/market").status_code == 200
    client.cookies.set("se_session", "9999999999.invalid")
    assert client.get("/api/market").status_code == 401


def test_failed_login_is_throttled(client):
    for _ in range(10):
        assert (
            client.post(
                "/api/session", json={"token": "wrong"}, headers=ORIGIN
            ).status_code
            == 401
        )
    assert (
        client.post("/api/session", json={"token": TOKEN}, headers=ORIGIN).status_code
        == 429
    )


def test_atomic_import_idempotence_and_no_forward_data(db):
    content = prices()
    assert store.import_csv(db, content, "test provider") == 25
    first = store.snapshot(db)[0]
    assert first["momentum"] == pytest.approx(124 / 104 - 1)
    assert first["signals"]["price"]["source"].startswith("Preis 20T:")
    store.import_csv(db, content, "test provider")
    with store.connect(db) as c:
        assert c.execute("select count(*) from observations").fetchone()[0] == 25
        original = c.execute(
            "select score from signals where date='2025-01-21'"
        ).fetchone()[0]
    # Adding future prices must never change a past score.
    store.import_csv(db, "symbol,date,close\nSAP.DE,2025-02-01,1000\n", "test provider")
    with store.connect(db) as c:
        assert (
            c.execute("select score from signals where date='2025-01-21'").fetchone()[0]
            == original
        )
    invalid = "symbol,date,close\nOTHER,2025-01-01,10\nOTHER,2025-01-02,NaN\n"
    with pytest.raises(ValueError):
        store.import_csv(db, invalid, "test provider")
    assert all(r["symbol"] != "OTHER" for r in store.snapshot(db))


def test_unknown_signal_rolls_back_whole_batch(db):
    store.import_csv(db, prices(), "test provider")
    with pytest.raises(ValueError):
        store.import_csv(
            db,
            "symbol,date,layer,score\nSAP.DE,2025-01-25,finbert,0.5\nNOPE,2025-01-25,llm,0.2\n",
            "model",
            "signals",
        )
    assert "finbert" not in store.snapshot(db)[0]["signals"]


def test_durable_watchlist_and_multi_asset_catalog(client, db):
    login(client)
    response = client.post("/api/import?source=CSV", content=prices(), headers=ORIGIN)
    assert response.status_code == 200
    response = client.put(
        "/api/watchlist/SAP.DE", json={"watched": True}, headers=ORIGIN
    )
    assert response.status_code == 200
    assert store.snapshot(db)[0]  # The result persists outside the HTTP request.
    with store.connect(db) as c:
        assert c.execute("select symbol from watchlist").fetchone()[0] == "SAP.DE"
    market = client.get("/api/market").json()["securities"]
    catalog = {r["symbol"]: r for r in market}
    assert catalog["CC=F"]["asset_class"] == "future"
    assert catalog["CC=F"]["close"] is None and catalog["CC=F"]["date"] is None
    assert catalog["SGLN.L"]["asset_class"] == "etc"
    assert catalog["SPY"]["asset_class"] == "etf"
    assert len(catalog["CC=F"]["exposures"]) == 2
    assert client.get("/api/securities/CC=F").json()["history"] == []
    assert client.get("/api/securities/DOESNOTEXIST").status_code == 404


def test_upload_limits_and_validation(client):
    login(client)
    assert (
        client.post(
            "/api/import?source=x",
            content=b"x" * (10 * 1024 * 1024 + 1),
            headers=ORIGIN,
        ).status_code
        == 413
    )
    assert (
        client.post(
            "/api/import?source=x",
            content="symbol,date,close\nX,2099-01-01,10",
            headers=ORIGIN,
        ).status_code
        == 422
    )
    assert (
        client.post(
            "/api/import?source=x&kind=bad", content="x", headers=ORIGIN
        ).status_code
        == 422
    )


def test_thesis_adapter_ignores_forward_returns_and_does_not_fabricate_prices(
    db, tmp_path
):
    source = tmp_path / "thesis.sqlite"
    with sqlite3.connect(source) as c:
        c.execute(
            "create table s1_price_panel_scored(yahoo_ticker,date,company,country,sector,mom_12_1,vol_60d,s1_price_score,ret_fwd_1d)"
        )
        c.execute(
            "insert into s1_price_panel_scored values(?,?,?,?,?,?,?,?,?)",
            ("SAP.DE", "2025-01-01", "SAP", "DE", "Software", 0.1, 0.2, 0.7, 999),
        )
    assert store.import_thesis(db, source) == 1
    item = store.snapshot(db)[0]
    assert item["close"] is None
    assert item["momentum"] == 0.1
    assert item["signals"]["price"]["score"] == 0.7
    assert "ret_fwd_1d" not in item
    with pytest.raises(ValueError):
        store.import_thesis(db, tmp_path / "missing")


def test_weather_metrics_missing_days_and_feed_failures(db, monkeypatch):
    assert context.longest_dry_spell([0, 0, None, 0, 0.5, 0, 2]) == 3

    def broken(*args):
        raise ValueError("private provider response")

    monkeypatch.setattr(context, "weather", broken)
    result = context.sync(db, "weather")
    assert result["status"] == "error"
    assert "private" not in result["message"]
    with pytest.raises(ValueError, match="15 Minuten"):
        context.sync(db, "weather")


def test_weather_feed_aggregation_and_provenance(db, monkeypatch):
    end = date.today() - timedelta(days=7)
    days = [(end - timedelta(days=29 - i)).isoformat() for i in range(30)]
    monkeypatch.setattr(
        context,
        "fetch_json",
        lambda *args: {
            "daily": {
                "time": days,
                "precipitation_sum": [0.0] * 5 + [2.0] * 25,
                "temperature_2m_max": [36.0] * 10 + [30.0] * 20,
            }
        },
    )
    assert context.weather(db) == 120
    result = context.snapshot(db)
    gh = {
        r["metric"]: r["value"] for r in result["series"] if r["region"] == "GH-ASHANTI"
    }
    assert gh == {"rain_30d": 50, "dry_spell_30d": 5, "heat_days_30d": 10}
    assert all(r["source"].startswith("Open-Meteo") for r in result["series"])


def test_event_urls_and_deduplication(db, monkeypatch):
    with pytest.raises(ValueError):
        context.safe_url("javascript:alert(1)")
    monkeypatch.setattr(
        context,
        "fetch_json",
        lambda *args: {
            "articles": [
                {
                    "url": "https://example.org/article",
                    "title": "Cocoa weather report",
                    "seendate": "20251001T100000Z",
                }
            ]
        },
    )
    context.news(db)
    context.news(db)
    assert len(context.snapshot(db)["events"]) == 2


def workbook_bytes():
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Monthly Prices"
    sheet.cell(5, 2, "Gold")
    sheet.cell(5, 3, "Cocoa")
    sheet.cell(6, 2, "($/troy oz)")
    sheet.cell(6, 3, "($/kg)")
    for i in range(14):
        year, month = 2023 + i // 12, i % 12 + 1
        (
            sheet.append([f"{year}M{month:02}", 100 + i, 2 + i / 100])
            if sheet.max_row >= 6
            else None
        )
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_commodity_reference_preserves_units_and_monthly_frequency(db):
    content = context.parse_worldbank(workbook_bytes())
    assert store.import_csv(db, content, "World Bank") == 28
    items = {r["symbol"]: r for r in store.snapshot(db)}
    assert items["COCOA-WB"]["unit"] == "kg"
    assert items["GOLD-WB"]["unit"] == "troy_ounce"
    assert items["COCOA-WB"]["frequency"] == "monthly"
    assert items["COCOA-WB"]["momentum"] == pytest.approx((2.13 / 2.01) - 1)
    assert "12M" in items["COCOA-WB"]["metric_basis"]
    with pytest.raises(ValueError, match="Frequenzen"):
        store.import_csv(db, "symbol,date,close\nGOLD-WB,2025-01-01,200\n", "daily")
    assert items["GOLD-WB"]["close"] == 113


def test_auth_payload_bound_before_json_parsing(client):
    response = client.post("/api/session", content=b"x" * 4097, headers=ORIGIN)
    assert response.status_code == 413


def test_monthly_gaps_do_not_become_twelve_month_signals(db):
    content = context.parse_worldbank(workbook_bytes())
    content = (
        "\n".join(line for line in content.splitlines() if "2023-02-28" not in line)
        + "\n"
    )
    store.import_csv(db, content, "World Bank")
    for item in store.snapshot(db):
        assert item["momentum"] is None
        assert "price" not in item["signals"]


def test_research_database_cannot_be_used_as_runtime_database(tmp_path):
    path = tmp_path / "original.sqlite"
    with sqlite3.connect(path) as c:
        c.execute("create table research(value integer)")
        c.execute("insert into research values(42)")
    original = path.read_bytes()
    with pytest.raises(RuntimeError, match="non-application"):
        store.initialize(path)
    assert path.read_bytes() == original
    with pytest.raises(ValueError, match="getrennt"):
        store.import_thesis(path, path)


def test_incomplete_csv_rows_return_validation_error(client):
    login(client)
    assert (
        client.post(
            "/api/import?source=x",
            content="symbol,date,close\nSAP.DE\n",
            headers=ORIGIN,
        ).status_code
        == 422
    )


def test_non_ascii_login_is_rejected_cleanly(client):
    assert (
        client.post(
            "/api/session", json={"token": "ä" * 40}, headers=ORIGIN
        ).status_code
        == 401
    )
