"""Explicit exposure links and bounded adapters; no implicit price forecasts."""

from __future__ import annotations

import hashlib
import calendar
import csv
import io
import re
import zipfile

import openpyxl
import json
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlparse

import httpx

from . import store

REGIONS = {
    "GH-ASHANTI": {"name": "Ashanti · Ghana", "lat": 6.7, "lon": -1.6},
    "CI-SOUBRE": {"name": "Soubré · Côte d’Ivoire", "lat": 5.79, "lon": -6.61},
}
CATALOG = [
    (
        "GOLD-WB",
        "Gold · World Bank Monatsreferenz",
        "",
        "Edelmetalle",
        "USD",
        "commodity_reference",
        "gold",
        "troy_ounce",
        "https://www.worldbank.org/en/research/commodity-markets",
    ),
    (
        "COCOA-WB",
        "Kakao · World Bank Monatsreferenz",
        "",
        "Agrarrohstoffe",
        "USD",
        "commodity_reference",
        "cocoa",
        "kg",
        "https://www.worldbank.org/en/research/commodity-markets",
    ),
    (
        "XAUUSD",
        "Gold Spot",
        "",
        "Edelmetalle",
        "USD",
        "commodity_spot",
        "gold",
        "troy_ounce",
        "https://www.lbma.org.uk/prices-and-data/precious-metal-prices",
    ),
    (
        "CC=F",
        "Kakao · kontinuierliche Futures-Reihe",
        "",
        "Agrarrohstoffe",
        "USD",
        "future",
        "cocoa",
        "metric_ton",
        "https://www.ice.com/products/7/Cocoa-Futures",
    ),
    (
        "SGLN.L",
        "iShares Physical Gold ETC",
        "United Kingdom",
        "Edelmetalle",
        "GBp",
        "etc",
        "gold",
        "security",
        "https://www.ishares.com/uk/individual/en/products/258441/ishares-physical-gold-etc-fund",
    ),
    (
        "SPY",
        "SPDR S&P 500 ETF Trust",
        "United States",
        "Breiter Aktienmarkt",
        "USD",
        "etf",
        "us_equity",
        "share",
        "https://www.ssga.com/us/en/individual/etfs/spdr-sp-500-etf-trust-spy",
    ),
]


def initialize_catalog(path):
    with store.connect(path) as c:
        for row in CATALOG:
            c.execute(
                "INSERT OR IGNORE INTO securities(symbol,name,country,sector,currency,asset_class,underlying,unit,reference_url) VALUES(?,?,?,?,?,?,?,?,?)",
                row,
            )
        for region in REGIONS:
            c.execute(
                "INSERT OR IGNORE INTO exposures VALUES(?,?,?,?,?)",
                (
                    "COCOA-WB",
                    "cocoa",
                    region,
                    "Anbaugebiet · repräsentativer Wetterpunkt, keine Flächengewichtung",
                    "https://doi.org/10.1371/journal.pclm.0000970",
                ),
            )
            c.execute(
                "INSERT OR IGNORE INTO exposures VALUES(?,?,?,?,?)",
                (
                    "CC=F",
                    "cocoa",
                    region,
                    "Anbaugebiet · repräsentativer Wetterpunkt, keine Flächengewichtung",
                    "https://doi.org/10.1371/journal.pclm.0000970",
                ),
            )
        for symbol in ["XAUUSD", "SGLN.L", "GOLD-WB"]:
            c.execute(
                "INSERT OR IGNORE INTO exposures VALUES(?,?,?,?,?)",
                (
                    symbol,
                    "gold",
                    "GLOBAL",
                    "Geopolitische Nachrichten · kontextabhängiger Zusammenhang",
                    "https://doi.org/10.1111/j.1540-6288.2010.00244.x",
                ),
            )


def safe_url(value):
    url = urlparse(value)
    if (
        url.scheme not in {"http", "https"}
        or not url.hostname
        or url.username
        or url.password
    ):
        raise ValueError("Ungültige Quellen-URL")
    return value[:2048]


def fetch_bytes(url, params=None, limit=2 * 1024 * 1024):
    # Fixed caller-owned endpoints only; never fetch URLs from articles or CSVs.
    with httpx.Client(timeout=20, follow_redirects=False) as client:
        with client.stream("GET", url, params=params) as response:
            response.raise_for_status()
            data = bytearray()
            for chunk in response.iter_bytes():
                data.extend(chunk)
                if len(data) > limit:
                    raise ValueError("Anbieterantwort zu groß")
    return bytes(data)


def fetch_json(url, params):
    return json.loads(fetch_bytes(url, params))


def longest_dry_spell(rain):
    longest = current = 0
    for value in rain:
        # Missing days break a run; they are not dry days.
        current = current + 1 if value is not None and value < 1 else 0
        longest = max(longest, current)
    return longest


def weather(path):
    end = date.today() - timedelta(days=7)  # ERA5 delay: avoid partial recent days.
    start = end - timedelta(days=29)
    count = 0
    for region, point in REGIONS.items():
        data = fetch_json(
            "https://archive-api.open-meteo.com/v1/archive",
            {
                "latitude": point["lat"],
                "longitude": point["lon"],
                "start_date": start.isoformat(),
                "end_date": end.isoformat(),
                "daily": "precipitation_sum,temperature_2m_max",
                "timezone": "GMT",
                "models": "era5",
            },
        )
        daily = data["daily"]
        days, rain, temp = (
            daily["time"],
            daily["precipitation_sum"],
            daily["temperature_2m_max"],
        )
        if (
            len(days) != 30
            or len(rain) != 30
            or len(temp) != 30
            or days[0] != start.isoformat()
            or days[-1] != end.isoformat()
        ):
            raise ValueError("Unvollständiges Wetterfenster")
        validated = []
        for day, r, t in zip(days, rain, temp):
            day = store.validate_date(day)
            r, t = store.finite(r, True), store.finite(t, True)
            if r is not None and r < 0:
                raise ValueError("Negativer Niederschlag")
            validated.append((day, r, t))
        if any(r is None or t is None for _, r, t in validated):
            raise ValueError("Fehlende Wetterwerte; vorheriger Stand bleibt erhalten")
        stamp = store.now()
        source = "Open-Meteo / ERA5 · Reanalyse"
        with store.connect(path) as c:
            for day, r, t in validated:
                for metric, value, unit in [
                    ("precipitation", r, "mm"),
                    ("temperature_max", t, "°C"),
                ]:
                    c.execute(
                        "INSERT INTO context_series VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(driver,region,date,metric) DO UPDATE SET value=excluded.value,observed_at=excluded.observed_at",
                        ("cocoa", region, day, metric, value, unit, source, stamp),
                    )
                    count += 1
            summaries = [
                ("rain_30d", sum(rain), "mm"),
                ("dry_spell_30d", longest_dry_spell(rain), "days"),
                ("heat_days_30d", sum(t >= 35 for t in temp), "days"),
            ]
            for metric, value, unit in summaries:
                c.execute(
                    "INSERT OR REPLACE INTO context_series VALUES(?,?,?,?,?,?,?,?)",
                    (
                        "cocoa",
                        region,
                        end.isoformat(),
                        metric,
                        value,
                        unit,
                        source,
                        stamp,
                    ),
                )
    return count


def news(path):
    count = 0
    queries = {
        "cocoa": '(cocoa OR cacao) (Ghana OR "Ivory Coast" OR "Cote d Ivoire") (drought OR rainfall OR conflict OR war OR disruption)',
        "gold": 'gold (war OR conflict OR sanctions OR "geopolitical risk")',
    }
    for driver, query in queries.items():
        data = fetch_json(
            "https://api.gdeltproject.org/api/v2/doc/doc",
            {
                "query": query,
                "mode": "artlist",
                "format": "json",
                "maxrecords": 30,
                "timespan": "7d",
                "sort": "datedesc",
            },
        )
        if "articles" not in data:
            raise ValueError("GDELT lieferte keine gültige Artikelantwort")
        records = []
        for item in data["articles"][:30]:
            url = safe_url(item["url"])
            seen = (
                datetime.strptime(item["seendate"], "%Y%m%dT%H%M%SZ")
                .replace(tzinfo=timezone.utc)
                .isoformat()
            )
            records.append(
                (
                    hashlib.sha256((driver + url).encode()).hexdigest(),
                    driver,
                    "GLOBAL",
                    seen,
                    store.text(item["title"], 1000),
                    url,
                    "GDELT DOC · unbestätigte Medienmeldung",
                    store.now(),
                )
            )
        with store.connect(path) as c:
            c.executemany(
                "INSERT OR IGNORE INTO events VALUES(?,?,?,?,?,?,?,?)", records
            )
        count += len(records)
    return count


def parse_worldbank(payload):
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
            raise ValueError("Workbook zu groß")
    workbook = openpyxl.load_workbook(
        io.BytesIO(payload), read_only=True, data_only=True
    )
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(
        [
            "symbol",
            "date",
            "close",
            "name",
            "currency",
            "asset_class",
            "underlying",
            "unit",
            "frequency",
        ]
    )
    try:
        sheet = workbook["Monthly Prices"]
        header = next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))
        units = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
        indices = {name: header.index(name) for name in ["Gold", "Cocoa"]}
        if (
            units[indices["Gold"]] != "($/troy oz)"
            or units[indices["Cocoa"]] != "($/kg)"
        ):
            raise ValueError("Unerwartete Rohstoffeinheit")
        cutoff = date.today().replace(day=1)
        count = 0
        for row in sheet.iter_rows(
            min_row=7, max_row=min(sheet.max_row, 1300), values_only=True
        ):
            if not isinstance(row[0], str) or not re.fullmatch(r"\d{4}M\d{2}", row[0]):
                continue
            year, month = int(row[0][:4]), int(row[0][5:])
            day = date(year, month, calendar.monthrange(year, month)[1])
            if day >= cutoff:
                continue  # Completed monthly averages only.
            for name, symbol, underlying, unit in [
                ("Gold", "GOLD-WB", "gold", "troy_ounce"),
                ("Cocoa", "COCOA-WB", "cocoa", "kg"),
            ]:
                value = row[indices[name]]
                if value in {None, "…", "..."}:
                    continue
                value = store.finite(value)
                if value <= 0:
                    raise ValueError("Ungültiger Referenzpreis")
                writer.writerow(
                    [
                        symbol,
                        day.isoformat(),
                        value,
                        f"{name} · World Bank Monatsreferenz",
                        "USD",
                        "commodity_reference",
                        underlying,
                        unit,
                        "monthly",
                    ]
                )
                count += 1
        if not count:
            raise ValueError("Keine Rohstoffbeobachtungen gefunden")
    finally:
        workbook.close()
    return out.getvalue()


def commodities(path):
    # Discover the current official workbook on a fixed origin. Validate the exact
    # destination pattern before fetching; no arbitrary redirects or URLs.
    html = fetch_bytes(
        "https://www.worldbank.org/en/research/commodity-markets"
    ).decode("utf-8")
    candidates = re.findall(
        r"https://thedocs\.worldbank\.org/en/doc/[a-f0-9]+-[0-9]+/related/CMO-Historical-Data-Monthly\.xlsx",
        html,
    )
    if not candidates:
        raise ValueError("Offizielle Workbook-URL nicht gefunden")
    payload = fetch_bytes(candidates[0], limit=8 * 1024 * 1024)
    content = parse_worldbank(payload)
    return store.import_csv(
        path, content, "World Bank Pink Sheet · Monatsmittel (nominal USD)"
    )


def sync(path, feed):
    if feed not in {"weather", "news", "commodities"}:
        raise ValueError("Unbekannter Feed")
    stamp = store.now()
    with store.connect(path) as c:
        c.execute("BEGIN IMMEDIATE")
        previous = c.execute("SELECT * FROM feed_runs WHERE feed=?", (feed,)).fetchone()
        if (
            previous
            and (
                datetime.now(timezone.utc)
                - datetime.fromisoformat(previous["last_attempt"])
            ).total_seconds()
            < 900
        ):
            raise ValueError("Feed wurde bereits versucht. Erneut nach 15 Minuten.")
        c.execute(
            "INSERT INTO feed_runs VALUES(?,?,NULL,'running','') ON CONFLICT(feed) DO UPDATE SET last_attempt=excluded.last_attempt,status='running',message=''",
            (feed, stamp),
        )
    try:
        count = {"weather": weather, "news": news, "commodities": commodities}[feed](
            path
        )
    except (
        httpx.HTTPError,
        ValueError,
        KeyError,
        TypeError,
        zipfile.BadZipFile,
    ) as exc:
        # No provider body or credential data in logs / UI.
        message = f"Quelle nicht verfügbar ({type(exc).__name__}); vorhandene Daten bleiben erhalten."
        with store.connect(path) as c:
            c.execute(
                "UPDATE feed_runs SET status='error',message=? WHERE feed=?",
                (message, feed),
            )
        return {"feed": feed, "status": "error", "message": message}
    with store.connect(path) as c:
        c.execute(
            "UPDATE feed_runs SET last_success=?,status='ok',message=? WHERE feed=?",
            (store.now(), f"{count} Beobachtungen verarbeitet", feed),
        )
    return {"feed": feed, "status": "ok", "rows": count}


def snapshot(path):
    with store.connect(path) as c:
        series = [
            dict(r)
            for r in c.execute(
                "SELECT * FROM context_series WHERE metric IN ('rain_30d','dry_spell_30d','heat_days_30d') AND date=(SELECT max(s.date) FROM context_series s WHERE s.region=context_series.region AND s.metric=context_series.metric)"
            )
        ]
        events = [
            dict(r)
            for r in c.execute("SELECT * FROM events ORDER BY date DESC LIMIT 60")
        ]
        feeds = [dict(r) for r in c.execute("SELECT * FROM feed_runs")]
    return {
        "regions": REGIONS,
        "series": series,
        "events": events,
        "feeds": feeds,
        "note": "Regionale Kontextdaten; keine Kausalitäts- oder Preisprognose. Trockenperioden sind kein SPI/SPEI-Dürreindex.",
    }
